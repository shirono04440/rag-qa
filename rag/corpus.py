# -*- coding: utf-8 -*-
"""语料加载：把目录下的文本类文件读成文档，供后续切分与建索引。

支持纯文本（md / txt / py / json / html / css / js / yaml）与 Word（.docx）。
PDF 需要额外依赖，见 README「接入其他数据」一节。
"""

from __future__ import annotations

import os
from dataclasses import dataclass
from typing import Iterable, List

TEXT_EXT = {".md", ".txt", ".py", ".json", ".html", ".css", ".js", ".yaml", ".yml"}
DOCX_EXT = {".docx"}
SKIP_DIRS = {
    ".git", ".venv", "venv", "env", "__pycache__", "build", "dist",
    "node_modules", ".idea", ".vscode", ".claude", "data",
}


@dataclass
class Doc:
    path: str          # 相对路径，作为引用来源
    text: str


def _iter_files(root: str) -> Iterable[str]:
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS]
        for name in filenames:
            ext = os.path.splitext(name)[1].lower()
            if ext in TEXT_EXT or ext in DOCX_EXT:
                yield os.path.join(dirpath, name)


def _read_docx(path: str) -> str:
    """用 python-docx 抽取段落与表格文本；未安装时返回空串。"""
    try:
        from docx import Document
    except ImportError:
        return ""
    doc = Document(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def load_documents(root: str, max_file_bytes: int = 2_000_000) -> List[Doc]:
    """读取 root 下所有支持的文档，返回 Doc 列表（跳过超大文件与空文件）。"""
    root = os.path.abspath(root)
    docs: List[Doc] = []
    for full in _iter_files(root):
        ext = os.path.splitext(full)[1].lower()
        try:
            if os.path.getsize(full) > max_file_bytes:
                continue
            if ext in DOCX_EXT:
                text = _read_docx(full)
            else:
                with open(full, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
        except (OSError, ValueError, KeyError):
            # 部分 docx 结构异常，python-docx 会抛 KeyError
            continue
        if not text.strip():
            continue
        rel = os.path.relpath(full, root).replace("\\", "/")
        docs.append(Doc(path=rel, text=text))
    docs.sort(key=lambda d: d.path)
    return docs
